"""Alumnos en riesgo y las vistas de reportes: acceso, descargas y lo que recuerdan."""

import datetime
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.courses.tests.factories import CourseFactory, EnrollmentFactory
from apps.courses.services.courses import archive_course
from apps.grades.tests.factories import ActivityFactory
from apps.grades.models import Activity
from apps.grades.services.gradebook import grade_activity
from apps.grades.services.schemes import ensure_scheme
from apps.grades.services.submissions import submit

from apps.reports.services import risk

pytestmark = pytest.mark.django_db

AHORA = timezone.now()
AYER = AHORA - datetime.timedelta(days=1)


def inscribir(course, apellido="Perez", nombre="Ana", matricula=""):
    return EnrollmentFactory(
        course=course, student__user=None, student__first_name=nombre, student__last_name=apellido,
        student__student_number=matricula,
    ).student


def evaluar(course, alumno, nota, **extra):
    """Califica a ``alumno`` con ``nota`` sobre 10 en una actividad nueva del grupo."""
    return grade_activity(ActivityFactory(course=course, **extra), alumno, score=nota)


def nivel(course, alumno):
    return next(r for r in risk.assess_course(course) if r.student == alumno)


# ── Riesgo ────────────────────────────────────────────────────────────

class TestRiesgo:
    @pytest.mark.parametrize("nota, esperado", [
        (10, risk.OK), (7, risk.OK), (6.99, risk.AT_RISK), (6, risk.AT_RISK), (5.99, risk.FAILING), (0, risk.FAILING),
    ])
    def test_niveles_por_promedio(self, course, nota, esperado):
        alumno = inscribir(course)
        evaluar(course, alumno, nota)
        assert nivel(course, alumno).level == esperado

    def test_el_margen_sigue_a_la_escala(self, course):
        esquema = ensure_scheme(course)
        esquema.scale_max, esquema.pass_mark = 100, 60
        esquema.save()
        alumno = inscribir(course)
        evaluar(course, alumno, 6.5, max_score=10)  # 65 sobre 100: por debajo de 70
        assert nivel(course, alumno).level == risk.AT_RISK

    def test_sin_calificaciones_ni_faltantes_es_sin_datos(self, course):
        alumno = inscribir(course)
        ActivityFactory(course=course)
        assert nivel(course, alumno).level == risk.NO_DATA

    def test_motivo_de_promedio(self, course):
        alumno = inscribir(course)
        evaluar(course, alumno, 4)
        assert "por debajo del mínimo de 6" in nivel(course, alumno).reasons[0]

    def test_faltantes_vencidos_ponen_en_riesgo(self, course):
        alumno = inscribir(course)
        for i in range(2):
            ActivityFactory(course=course, type=Activity.Type.HOMEWORK, title=f"Tarea {i}", due_at=AYER)
        resultado = nivel(course, alumno)
        assert resultado.level == risk.AT_RISK and resultado.missing == ["Tarea 0", "Tarea 1"]
        assert "2 actividades vencidas" in resultado.reasons[0]

    def test_un_faltante_solo_no_basta(self, course):
        alumno = inscribir(course)
        ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER)
        assert nivel(course, alumno).level == risk.NO_DATA

    def test_entregar_o_estar_calificado_no_es_faltar(self, course):
        alumno = inscribir(course)
        entregada = ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER)
        calificada = ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER)
        pendiente = ActivityFactory(course=course, type=Activity.Type.HOMEWORK, title="Pendiente", due_at=AYER)
        submit(entregada, alumno, text="x", now=AHORA)
        grade_activity(calificada, alumno, score=9)
        resultado = nivel(course, alumno)
        assert resultado.missing == ["Pendiente"] and pendiente

    def test_lo_que_aun_no_vence_o_es_borrador_no_cuenta(self, course):
        alumno = inscribir(course)
        ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AHORA + datetime.timedelta(days=3))
        ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER, published=False)
        ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER, published=False)
        assert nivel(course, alumno).missing == []

    def test_reprobado_con_faltantes_conserva_el_nivel_y_suma_motivos(self, course):
        alumno = inscribir(course)
        evaluar(course, alumno, 3)
        for _ in range(2):
            ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER)
        resultado = nivel(course, alumno)
        assert resultado.level == risk.FAILING and len(resultado.reasons) == 2

    def test_bien_con_faltantes_pasa_a_riesgo(self, course):
        alumno = inscribir(course)
        evaluar(course, alumno, 10)
        for _ in range(2):
            ActivityFactory(course=course, type=Activity.Type.HOMEWORK, due_at=AYER)
        assert nivel(course, alumno).level == risk.AT_RISK

    def test_orden_primero_los_peores(self, course):
        notas = {"Bien": 9, "Mal": 2, "Peor": 0, "Justo": 6.5}
        for apellido, nota in notas.items():
            evaluar(course, inscribir(course, apellido), nota)
        assert [r.student.last_name for r in risk.assess_course(course)] == ["Peor", "Mal", "Justo", "Bien"]

    def test_conteos(self, course):
        for apellido, nota in {"A": 9, "B": 2, "C": 6.5}.items():
            evaluar(course, inscribir(course, apellido), nota)
        inscribir(course, "D")
        assert risk.risk_counts(risk.assess_course(course)) == {risk.FAILING: 1, risk.AT_RISK: 1, risk.OK: 1, risk.NO_DATA: 1}

    def test_los_alumnos_de_otro_grupo_no_se_mezclan(self, course):
        alumno = inscribir(course)
        otro = CourseFactory(professor=course.professor)
        EnrollmentFactory(course=otro, student=alumno)
        evaluar(otro, alumno, 1)
        assert nivel(course, alumno).level == risk.NO_DATA


# ── Vistas ────────────────────────────────────────────────────────────

@pytest.fixture
def grupo(course):
    ana = inscribir(course, "Zapata", "Ana", "00123")
    evaluar(course, ana, 9)
    return ana


@pytest.mark.usefixtures('plan_docente')  # elegir y ordenar columnas es del plan Docente (ver test_planes_y_limites)
class TestDescargas:
    def test_boletas_del_grupo(self, professor_client, course, grupo):
        response = professor_client.get(reverse("reports:transcripts", args=[course.pk]))
        assert response.status_code == 200 and response["Content-Type"] == "application/pdf"
        assert response.content.startswith(b"%PDF-")
        assert response["Content-Disposition"].startswith('attachment; filename="boletas-')
        assert response["Cache-Control"] == "private, no-store" and response["X-Content-Type-Options"] == "nosniff"

    def test_boleta_de_un_alumno(self, professor_client, course, grupo):
        response = professor_client.get(reverse("reports:transcript", args=[course.pk, grupo.pk]))
        assert response.status_code == 200 and response.content.startswith(b"%PDF-")
        assert "boleta-00123-" in response["Content-Disposition"] and response["Content-Disposition"].endswith('.pdf"')

    def test_boleta_de_alguien_que_no_es_del_grupo(self, professor_client, course, grupo):
        from apps.courses.tests.factories import StudentRecordFactory

        mismo_org = StudentRecordFactory(organization=course.organization, user=None)
        otro_org = EnrollmentFactory().student
        for ajeno in (mismo_org, otro_org):
            assert professor_client.get(reverse("reports:transcript", args=[course.pk, ajeno.pk])).status_code == 404

    def test_el_nombre_del_archivo_es_seguro(self, professor_client, course, grupo):
        course.code = 'A/../"B\r\nX: y'
        course.save()
        disposicion = professor_client.get(reverse("reports:transcripts", args=[course.pk]))["Content-Disposition"]
        assert "\n" not in disposicion and "\r" not in disposicion and "/" not in disposicion.split("filename=")[1]

    def test_concentrado_por_defecto(self, professor_client, course, grupo):
        response = professor_client.get(reverse("reports:workbook_download", args=[course.pk]))
        assert response.status_code == 200 and response["Content-Type"].endswith("spreadsheetml.sheet")
        hoja = load_workbook(BytesIO(response.content))["Concentrado"]
        assert hoja["A1"].value == "Matrícula" and hoja["A2"].value == "00123"

    def test_concentrado_con_columnas_elegidas(self, professor_client, course, grupo):
        response = professor_client.get(reverse("reports:workbook_download", args=[course.pk]), {"cols": ["final", "email", "last_name"]})
        hoja = load_workbook(BytesIO(response.content))["Concentrado"]
        assert [c.value for c in hoja[1]] == ["Calificación final", "Correo", "Apellidos"]

    def test_columnas_inexistentes_son_400(self, professor_client, course, grupo):
        response = professor_client.get(reverse("reports:workbook_download", args=[course.pk]), {"cols": ["x", "../y"]})
        assert response.status_code == 400

    def test_recuerda_la_ultima_eleccion_de_cada_grupo(self, professor_client, course, grupo, professor):
        professor_client.get(reverse("reports:workbook_download", args=[course.pk]), {"cols": ["final", "last_name"]})
        otro = CourseFactory(professor=professor)
        inscribir(otro)
        # Sin parámetros, vuelve a usar lo último de ESTE grupo; el otro grupo sigue con lo de siempre.
        hoja = load_workbook(BytesIO(professor_client.get(reverse("reports:workbook_download", args=[course.pk])).content))["Concentrado"]
        assert [c.value for c in hoja[1]] == ["Calificación final", "Apellidos"]
        hoja_otro = load_workbook(BytesIO(professor_client.get(reverse("reports:workbook_download", args=[otro.pk])).content))["Concentrado"]
        assert hoja_otro["A1"].value == "Matrícula"

    def test_la_pantalla_de_columnas_marca_lo_recordado_y_ordena(self, professor_client, course, grupo):
        professor_client.get(reverse("reports:workbook_download", args=[course.pk]), {"cols": ["final", "last_name"]})
        filas = professor_client.get(reverse("reports:workbook", args=[course.pk])).context["rows"]
        assert [r["column"].key for r in filas][:2] == ["final", "last_name"]
        assert [r["checked"] for r in filas][:3] == [True, True, False] and sum(r["checked"] for r in filas) == 2

    def test_pantalla_de_columnas_por_defecto(self, professor_client, course):
        filas = professor_client.get(reverse("reports:workbook", args=[course.pk])).context["rows"]
        assert [r["column"].key for r in filas if r["checked"]] == ["student_number", "last_name", "first_name", "activities", "final", "status"]

    def test_un_grupo_archivado_tambien_se_reporta(self, professor_client, course, grupo):
        archive_course(course)
        for nombre in ("reports:transcripts", "reports:workbook_download", "reports:risk", "reports:workbook"):
            assert professor_client.get(reverse(nombre, args=[course.pk])).status_code == 200, nombre


class TestPantallaDeRiesgo:
    def test_muestra_a_quien_necesita_atencion(self, professor_client, course):
        evaluar(course, inscribir(course, "Mal", "Beto"), 2)
        evaluar(course, inscribir(course, "Bien", "Ana"), 9)
        response = professor_client.get(reverse("reports:risk", args=[course.pk]))
        html = response.content.decode()
        assert response.status_code == 200 and "Beto" in html and "por debajo del mínimo" in html
        assert [a.student.last_name for a in response.context["attention"]] == ["Mal"]
        assert response.context["counts"][risk.OK] == 1
        assert reverse("reports:transcript", args=[course.pk, response.context["attention"][0].student.pk]) in html

    def test_nadie_en_riesgo(self, professor_client, course):
        evaluar(course, inscribir(course), 9)
        assert "Nadie necesita atención" in professor_client.get(reverse("reports:risk", args=[course.pk])).content.decode()

    def test_los_nombres_se_escapan(self, professor_client, course):
        evaluar(course, inscribir(course, "<script>x</script>"), 1)
        html = professor_client.get(reverse("reports:risk", args=[course.pk])).content.decode()
        assert "<script>x</script>" not in html


class TestAcceso:
    URLS = [
        ("reports:transcripts", lambda c, s: [c.pk]), ("reports:transcript", lambda c, s: [c.pk, s.pk]),
        ("reports:workbook", lambda c, s: [c.pk]), ("reports:workbook_download", lambda c, s: [c.pk]),
        ("reports:risk", lambda c, s: [c.pk]),
    ]

    @pytest.mark.parametrize("nombre, args", URLS)
    def test_el_alumno_no_accede(self, student_client, course, grupo, nombre, args):
        assert student_client.get(reverse(nombre, args=args(course, grupo))).status_code == 302

    @pytest.mark.parametrize("nombre, args", URLS)
    def test_anonimo_va_al_login(self, client, course, grupo, nombre, args):
        assert client.get(reverse(nombre, args=args(course, grupo))).url.startswith("/login/")

    @pytest.mark.parametrize("nombre, args", URLS)
    def test_el_grupo_de_otro_docente_es_404(self, professor_client, nombre, args):
        ajeno = EnrollmentFactory()
        assert professor_client.get(reverse(nombre, args=args(ajeno.course, ajeno.student))).status_code == 404

    @pytest.mark.parametrize("nombre, args", URLS)
    def test_el_grupo_de_otro_docente_de_la_misma_organizacion_es_404(self, professor_client, course, nombre, args):
        from apps.accounts.tests.factories import ProfessorFactory
        from apps.organizations.tests.factories import MembershipFactory
        from apps.organizations.models import Membership

        colega = ProfessorFactory()
        MembershipFactory(user=colega, organization=course.organization, role=Membership.Role.TEACHER)
        ajeno = EnrollmentFactory(course=CourseFactory(professor=colega, organization=course.organization))
        assert ajeno.course.organization == course.organization
        assert professor_client.get(reverse(nombre, args=args(ajeno.course, ajeno.student))).status_code == 404

    def test_enlaces_en_el_grupo_y_la_libreta(self, professor_client, course, grupo):
        detalle = professor_client.get(reverse("course_detail", args=[course.pk])).content.decode()
        libreta = professor_client.get(reverse("grades:gradesheet", args=[course.pk])).content.decode()
        for nombre in ("reports:risk", "reports:workbook", "reports:transcripts"):
            assert reverse(nombre, args=[course.pk]) in detalle
        assert reverse("reports:transcript", args=[course.pk, grupo.pk]) in detalle
        assert reverse("reports:risk", args=[course.pk]) in libreta and reverse("reports:workbook", args=[course.pk]) in libreta
