"""Concentrado en Excel: una fila por alumno, con las columnas que el docente elija y en el orden que elija.

El formato que pide cada sistema de control escolar varía, así que no se intenta adivinar uno: el docente
arma el suyo. La matrícula va siempre como texto (Excel se come los ceros a la izquierda de un número) y
nada de lo que venga de un alumno se interpreta como fórmula.
"""

from dataclasses import dataclass
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.grades.services.grading import ZERO, _dec
from apps.grades.services.gradesheet import fmt

from .transcript import build_transcripts


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    help: str = ''
    identity: bool = False  # datos del alumno: se dejan fijos al desplazarse


COLUMNS = (
    Column('student_number', 'Matrícula', identity=True),
    Column('last_name', 'Apellidos', identity=True),
    Column('first_name', 'Nombre(s)', identity=True),
    Column('email', 'Correo', identity=True),
    Column('activities', 'Calificación de cada actividad', 'Una columna por actividad publicada.'),
    Column('criteria', 'Promedio de cada criterio', 'Una columna por criterio del esquema, en la escala del grupo.'),
    Column('final', 'Calificación final'),
    Column('status', 'Resultado', 'Aprobado o No aprobado.'),
)
BY_KEY = {column.key: column for column in COLUMNS}
DEFAULT_KEYS = ('student_number', 'last_name', 'first_name', 'activities', 'final', 'status')

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HEADER_FILL = PatternFill('solid', fgColor='1E3A5F')
HEADER_FONT = Font(bold=True, color='FFFFFF')
MAX_WIDTH = 42


def parse_columns(keys):
    """Claves válidas y sin repetir, en el orden recibido. Lo desconocido se ignora."""
    seen = []
    for key in keys:
        if key in BY_KEY and key not in seen:
            seen.append(key)
    return seen


def _text(cell, value):
    """Escribe texto sin que Excel lo trate como fórmula, aunque empiece con =, +, - o @."""
    cell.value = str(value)
    cell.data_type = 's'


def _number_format(decimals):
    return '0' if decimals == 0 else '0.' + '0' * decimals


def build_workbook(course, keys):
    """Libro de Excel (bytes) con el concentrado del grupo. ``keys`` son las columnas elegidas, en orden."""
    keys = parse_columns(keys) or list(DEFAULT_KEYS)
    transcripts = build_transcripts(course)
    sample = transcripts[0] if transcripts else None

    # Las actividades y los criterios son los mismos para todos los alumnos: se leen de cualquier boleta.
    activity_headers, criterion_headers = [], []
    if sample is not None:
        for period in sample.periods:
            for criterion in period.criteria:
                label = f'{period.name} · {criterion.name}' if sample.multiple_periods else criterion.name
                criterion_headers.append(f'{label} ({criterion.weight} %)')
                activity_headers.extend(f'{line.title} (/{line.max_score})' for line in criterion.lines)

    headers = []
    for key in keys:
        if key == 'activities':
            headers.extend(activity_headers)
        elif key == 'criteria':
            headers.extend(criterion_headers)
        else:
            headers.append(BY_KEY[key].label)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Concentrado'
    sheet.append([None] * len(headers))
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        _text(cell, header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')

    decimals = sample.decimals if sample else 2
    final_format = _number_format(decimals)
    for transcript in transcripts:
        student = transcript.student
        values = []  # (valor, tipo, formato): tipo 'text' | 'number' | 'final'
        for key in keys:
            if key == 'student_number':
                values.append((student.student_number, 'text', '@'))  # texto: no se pierden los ceros a la izquierda
            elif key == 'last_name':
                values.append((student.last_name, 'text', None))
            elif key == 'first_name':
                values.append((student.first_name, 'text', None))
            elif key == 'email':
                values.append((student.email, 'text', None))
            elif key == 'activities':
                values.extend((line.score, 'number', '0.##') for line in transcript.criterion_lines())
            elif key == 'criteria':
                values.extend((c.value_number, 'number', final_format) for p in transcript.periods for c in p.criteria)
            elif key == 'final':
                values.append((transcript.final_number, 'number', final_format))
            elif key == 'status':
                values.append((transcript.status if transcript.passed is not None else '', 'text', None))
        sheet.append([None] * len(values))
        row = sheet.max_row
        for index, (value, kind, number_format) in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index)
            if number_format:
                cell.number_format = number_format
            if value is None or value == '':
                continue
            if kind == 'text':
                _text(cell, value)
            else:
                cell.value = float(value)

    for index in range(1, len(headers) + 1):
        letter = get_column_letter(index)
        longest = max((len(str(cell.value)) for cell in sheet[letter] if cell.value is not None), default=8)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 10), MAX_WIDTH)
    sheet.row_dimensions[1].height = 32

    fixed = 0
    for key in keys:
        if not BY_KEY[key].identity:
            break
        fixed += 1
    sheet.freeze_panes = sheet.cell(row=2, column=fixed + 1)

    _summary_sheet(workbook, course, transcripts)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary_sheet(workbook, course, transcripts):
    sheet = workbook.create_sheet('Resumen')
    finals = [_dec(t.final_number) for t in transcripts if t.final_number is not None]
    decimals = transcripts[0].decimals if transcripts else 2
    rows = [
        ('Grupo', f'{course.name} ({course.code})'),
        ('Ciclo escolar', str(course.academic_year)),
        ('Docente', course.professor.get_full_name() or course.professor.get_username()),
        ('Generado', timezone.localtime().strftime('%d/%m/%Y %H:%M')),
        ('Alumnos inscritos', len(transcripts)),
        ('Alumnos con calificación', len(finals)),
        ('Promedio del grupo', fmt(sum(finals, ZERO) / len(finals), decimals, trim=False) if finals else ''),
        ('Aprobados', sum(1 for t in transcripts if t.passed)),
        ('No aprobados', sum(1 for t in transcripts if t.passed is False)),
        ('Mínimo aprobatorio', transcripts[0].pass_mark if transcripts else ''),
        ('Escala', transcripts[0].scale_max if transcripts else ''),
    ]
    for label, value in rows:
        sheet.append([None, None])
        row = sheet.max_row
        _text(sheet.cell(row=row, column=1), label)
        sheet.cell(row=row, column=1).font = Font(bold=True)
        if isinstance(value, int):
            sheet.cell(row=row, column=2).value = value
        else:
            _text(sheet.cell(row=row, column=2), value)
        sheet.cell(row=row, column=2).alignment = Alignment(horizontal='left')
    sheet.column_dimensions['A'].width = 26
    sheet.column_dimensions['B'].width = 44
