"""Importación de la lista de un grupo desde un archivo CSV o Excel.

Tres pasos, todos sin tocar la base hasta ``apply_import``:

1. ``read_table``: lee el archivo como texto (la matrícula nunca se interpreta como número).
2. ``detect_mapping`` / ``analyze``: asocia columnas con campos y prevé qué pasaría con cada fila.
3. ``apply_import``: crea o actualiza alumnos y los inscribe en el grupo.

La matrícula es la llave de re-importación: subir la lista actualizada no duplica alumnos. Sin
matrícula, se busca por correo y, en último caso, por nombre entre los alumnos ya inscritos en el grupo.
"""

import csv
import io
import re
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from apps.organizations import limits

from ..models import Enrollment, Student
from .text import normalize_text

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000

# Campos importables, en el orden en que se muestran.
FIELDS = {
    'student_number': 'Matrícula',
    'first_name': 'Nombre(s)',
    'last_name': 'Apellidos',
    'email': 'Correo',
}
REQUIRED_FIELD = 'first_name'

# Encabezados reconocidos (normalizados: minúsculas, sin acentos ni signos).
ALIASES = {
    'student_number': {
        'matricula', 'clave', 'id', 'folio', 'control', 'no control', 'no de control', 'num control',
        'num de control', 'numero de control',
    },
    'first_name': {'nombre', 'nombres', 'nombre s', 'name', 'first name'},
    'last_name': {'apellido', 'apellidos', 'last name', 'apellido paterno'},
    'email': {'correo', 'email', 'e mail', 'correo electronico', 'mail'},
}

SCIENTIFIC = re.compile(r'\d+(\.\d+)?[eE][+-]?\d+')


class ImportFileError(ValueError):
    """El archivo no se puede leer; el mensaje se muestra tal cual al docente."""


# ── Lectura ───────────────────────────────────────────────────────────

def _cell_text(value, number_format='General'):
    """Texto de una celda de Excel. Un número entero conserva sus ceros si la celda tenía formato 00000."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        text = str(value)
        if re.fullmatch(r'0+', number_format or ''):
            text = text.zfill(len(number_format))
        return text
    return str(value).strip()


def _read_csv(data):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError('No se pudo leer el archivo: guárdalo como CSV UTF-8.')
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - archivo no confiable: openpyxl puede lanzar muchos tipos de error
        raise ImportFileError('No se pudo abrir el archivo de Excel. ¿Está dañado o protegido con contraseña?') from error
    try:
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows():
            rows.append([_cell_text(cell.value, getattr(cell, 'number_format', 'General')) for cell in row])
            if len(rows) > MAX_ROWS + 1:
                break
        return rows
    finally:
        workbook.close()


def read_table(filename, data):
    """(encabezados, filas) del archivo, todo como texto. Lanza ``ImportFileError`` si no se puede usar."""
    if len(data) > MAX_BYTES:
        raise ImportFileError('El archivo pesa más de 2 MB. Divide la lista en varios archivos.')
    name = (filename or '').lower()
    if name.endswith('.xlsx'):
        rows = _read_xlsx(data)
    elif name.endswith(('.csv', '.txt')):
        rows = _read_csv(data)
    elif name.endswith('.xls'):
        raise ImportFileError('El formato .xls antiguo no se admite: guárdalo como .xlsx o CSV.')
    else:
        raise ImportFileError('Formato no admitido. Sube un archivo .xlsx o .csv.')

    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise ImportFileError('El archivo está vacío.')
    headers, body = rows[0], rows[1:]
    if not body:
        raise ImportFileError('El archivo solo tiene encabezados: no hay alumnos que importar.')
    if len(body) > MAX_ROWS:
        raise ImportFileError(f'El archivo tiene más de {MAX_ROWS} alumnos. Divide la lista en varios archivos.')

    width = len(headers)
    headers = [header or f'Columna {number}' for number, header in enumerate(headers, start=1)]
    body = [(row + [''] * width)[:width] for row in body]
    return headers, body


# ── Mapeo ─────────────────────────────────────────────────────────────

def detect_mapping(headers):
    """{campo: índice de columna} según los encabezados reconocidos; los demás campos quedan en None."""
    mapping = {key: None for key in FIELDS}
    for index, header in enumerate(headers):
        normalized = normalize_text(header)
        for key, aliases in ALIASES.items():
            if mapping[key] is None and normalized in aliases:
                mapping[key] = index
                break
    return mapping


def validate_mapping(mapping, width):
    """Mensajes de error del mapeo (lista vacía si es válido)."""
    errors = []
    if mapping.get(REQUIRED_FIELD) is None:
        errors.append(f'Indica qué columna es "{FIELDS[REQUIRED_FIELD]}".')
    used = [index for index in mapping.values() if index is not None]
    if any(index >= width or index < 0 for index in used):
        errors.append('Hay una columna que no existe en el archivo.')
    if len(used) != len(set(used)):
        errors.append('Una misma columna no puede usarse para dos campos.')
    return errors


# ── Análisis ──────────────────────────────────────────────────────────

@dataclass
class RowPlan:
    line: int
    values: dict
    action: str = 'create'  # create | update | unchanged | error
    enroll: bool = True
    student: Student | None = None
    match: str = ''  # por qué se asoció a un alumno existente: matrícula, correo o nombre
    changes: list = field(default_factory=list)
    errors: list = field(default_factory=list)

    @property
    def changes_labels(self):
        return [FIELDS[key] for key in self.changes]


@dataclass
class ImportPlan:
    rows: list

    def count(self, action):
        return sum(1 for row in self.rows if row.action == action)

    # Para las plantillas, que no pueden pasar argumentos.
    @property
    def count_create(self):
        return self.count('create')

    @property
    def count_update(self):
        return self.count('update')

    @property
    def count_unchanged(self):
        return self.count('unchanged')

    @property
    def count_error(self):
        return self.count('error')

    @property
    def to_enroll(self):
        return sum(1 for row in self.rows if row.action != 'error' and row.enroll)


def _name_key(first, last):
    return (normalize_text(first), normalize_text(last))


def analyze(course, rows, mapping):
    """Prevé qué haría la importación con cada fila (sin escribir nada)."""
    organization = course.organization
    students = list(Student.objects.for_organization(organization))
    by_number = {s.student_number: s for s in students if s.student_number}
    by_email = {}
    for student in students:
        if student.email:
            by_email.setdefault(student.email.lower(), []).append(student)
    enrolled_ids = set(course.enrollments.values_list('student_id', flat=True))
    enrolled_by_name = {}
    for student in students:
        if student.pk in enrolled_ids:
            enrolled_by_name.setdefault(_name_key(student.first_name, student.last_name), []).append(student)

    plans, seen_numbers, claimed = [], {}, set()
    for index, row in enumerate(rows):
        line = index + 2  # la fila 1 son los encabezados
        values = {
            key: (row[column].strip() if column is not None else '') for key, column in mapping.items()
        }
        plan = RowPlan(line=line, values=values)
        plans.append(plan)
        number, email = values['student_number'], values['email']

        if not values['first_name']:
            plan.errors.append('Falta el nombre.')
        if len(number) > Student._meta.get_field('student_number').max_length:
            plan.errors.append('La matrícula es demasiado larga.')
        if SCIENTIFIC.fullmatch(number):
            plan.errors.append(
                'La matrícula parece estar en notación científica (Excel la convirtió a número). '
                'Da formato de texto a esa columna y vuelve a guardar.'
            )
        if email:
            try:
                validate_email(email)
            except ValidationError:
                plan.errors.append('El correo no es válido.')
        if number:
            if number in seen_numbers:
                plan.errors.append(f'La matrícula ya aparece en la fila {seen_numbers[number]}.')
            else:
                seen_numbers[number] = line
        if plan.errors:
            plan.action = 'error'
            continue

        student = _find_student(plan, by_number, by_email, enrolled_by_name)
        if plan.errors:
            plan.action = 'error'
            continue
        if student is not None and student.pk in claimed:
            plan.errors.append('Este alumno ya aparece en otra fila del archivo.')
            plan.action = 'error'
            continue

        if student is None:
            plan.action = 'create'
            plan.enroll = True
        else:
            claimed.add(student.pk)
            plan.student = student
            plan.enroll = student.pk not in enrolled_ids
            plan.changes = [
                key for key in FIELDS if values[key] and values[key] != getattr(student, key)
            ]
            plan.action = 'update' if plan.changes else 'unchanged'
    return ImportPlan(rows=plans)


def _find_student(plan, by_number, by_email, enrolled_by_name):
    """Alumno existente que corresponde a la fila: por matrícula, si no por correo, si no por nombre."""
    number, email = plan.values['student_number'], plan.values['email']

    if number and number in by_number:
        plan.match = 'matrícula'
        return by_number[number]

    if email:
        candidates = by_email.get(email.lower(), [])
        if len(candidates) > 1:
            plan.errors.append('Hay varios alumnos con ese correo; no se puede saber cuál es.')
            return None
        if candidates:
            candidate = candidates[0]
            if number and candidate.student_number and candidate.student_number != number:
                plan.errors.append(
                    f'El correo ya pertenece a {candidate}, que tiene la matrícula {candidate.student_number}.'
                )
                return None
            plan.match = 'correo'
            return candidate

    # Último recurso: mismo nombre entre los inscritos del grupo que aún no tienen matrícula (o no se da una).
    candidates = enrolled_by_name.get(_name_key(plan.values['first_name'], plan.values['last_name']), [])
    if number:
        candidates = [c for c in candidates if not c.student_number]
    if len(candidates) == 1:
        plan.match = 'nombre'
        return candidates[0]
    if len(candidates) > 1:
        plan.errors.append('Hay varios alumnos con ese nombre en el grupo; usa la matrícula o el correo.')
    return None


# ── Aplicación ────────────────────────────────────────────────────────

def new_active_students(course, plan):
    """Cuántos alumnos sumarían al tope del plan si se aplica la importación (los que ya cuentan no suman)."""
    enrolled = [row for row in plan.rows if row.action != 'error' and row.enroll]
    created = sum(1 for row in enrolled if row.action == 'create')
    existing = [row.student.pk for row in enrolled if row.action != 'create' and row.student is not None]
    return created + limits.new_students_among(course.organization, existing)


def check_plan_limit(course, plan):
    """Lanza ``PlanLimitError`` si la importación completa pasaría el tope de alumnos del plan; si no, nada."""
    if limits.plan_for(course.organization).students is None:
        return
    limits.check_limit(course.organization, 'students', adding=new_active_students(course, plan))


@transaction.atomic
def apply_import(course, plan):
    """Aplica las filas sin error: crea o actualiza alumnos y los inscribe. Devuelve un resumen.

    Si pasa el tope de alumnos del plan se rechaza completa (``PlanLimitError``), no se corta a la mitad.
    """
    check_plan_limit(course, plan)
    summary = {'created': 0, 'updated': 0, 'enrolled': 0, 'skipped': plan.count('error')}
    for row in plan.rows:
        if row.action == 'error':
            continue
        if row.action == 'create':
            student = Student.objects.create(organization=course.organization, **row.values)
            summary['created'] += 1
        else:
            student = row.student
            if row.action == 'update':
                for key in row.changes:
                    setattr(student, key, row.values[key])
                student.save()
                summary['updated'] += 1
        if row.enroll:
            Enrollment.objects.create(student=student, course=course)
            summary['enrolled'] += 1
    return summary
