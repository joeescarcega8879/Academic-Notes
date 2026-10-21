"""Extracción del texto de un temario subido (PDF con texto, Word moderno o texto plano).

Nada de esto usa IA ni servicios externos. Un PDF escaneado (solo imágenes) no tiene texto que leer: se devuelve
vacío y el docente pega el texto.
"""

import os
import re

# Topes para que un archivo hostil o enorme no agote la memoria ni el tiempo de la tarea.
MAX_PAGES = 300
MAX_CHARS = 200_000


class ExtractionError(Exception):
    """El archivo no se pudo leer (dañado, protegido con contraseña, formato inesperado)."""


def clean_text(text):
    """Quita nulos y espacios sobrantes, une saltos de línea de más y recorta al tope."""
    text = text.replace('\x00', '').replace('\r\n', '\n').replace('\r', '\n').replace(' ', ' ')
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n')]
    text = '\n'.join(lines)
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    return text[:MAX_CHARS]


def _from_pdf(file):
    from pypdf import PdfReader

    try:
        reader = PdfReader(file)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ExtractionError('El PDF está protegido con contraseña.')
        parts, size = [], 0
        for page in reader.pages[:MAX_PAGES]:
            chunk = page.extract_text() or ''
            parts.append(chunk)
            size += len(chunk)
            if size >= MAX_CHARS:
                break
    except ExtractionError:
        raise
    except Exception as error:  # noqa: BLE001 - un archivo hostil puede romper pypdf de muchas formas
        raise ExtractionError('No se pudo leer el PDF.') from error
    return '\n\n'.join(parts)


def _from_docx(file):
    from docx import Document

    try:
        document = Document(file)
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(' | '.join(cell.text.strip() for cell in row.cells))
    except Exception as error:  # noqa: BLE001 - zip dañado, XML mal formado, partes que faltan…
        raise ExtractionError('No se pudo leer el documento de Word.') from error
    return '\n'.join(lines)


def _from_txt(file):
    raw = file.read(MAX_CHARS * 4)
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode('latin-1')


READERS = {'.pdf': _from_pdf, '.docx': _from_docx, '.txt': _from_txt}


def extract_text(file, filename):
    """Texto limpio de ``file`` (un objeto de archivo abierto en binario) según la extensión de ``filename``.

    Devuelve ``''`` si el archivo es válido pero no trae texto (p. ej. un PDF escaneado). Lanza
    ``ExtractionError`` si no se puede leer.
    """
    reader = READERS.get(os.path.splitext(filename)[1].lower())
    if reader is None:
        raise ExtractionError('Tipo de archivo no admitido.')
    return clean_text(reader(file))
