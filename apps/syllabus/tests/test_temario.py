"""Temario y semanas: extracción de texto, guardado, operaciones sobre semanas, detección y fechas."""

import datetime
import io

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from reportlab.pdfgen import canvas

from apps.courses.tests.factories import CourseFactory
from apps.syllabus.models import MAX_SYLLABUS_BYTES, Syllabus, SyllabusWeek, validate_syllabus_file
from apps.syllabus.services import extraction, text as texto, weeks as semanas
from apps.syllabus.tasks import extract_syllabus_text
from apps.syllabus.tests.factories import SyllabusFactory, SyllabusWeekFactory

pytestmark = pytest.mark.django_db


def pdf_con(*paginas):
    buffer = io.BytesIO()
    lienzo = canvas.Canvas(buffer)
    for pagina in paginas:
        for i, linea in enumerate(pagina):
            lienzo.drawString(72, 750 - 18 * i, linea)
        lienzo.showPage()
    lienzo.save()
    return buffer.getvalue()


def docx_con(*parrafos, tabla=None):
    documento = Document()
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    if tabla:
        t = documento.add_table(rows=len(tabla), cols=len(tabla[0]))
        for i, fila in enumerate(tabla):
            for j, celda in enumerate(fila):
                t.cell(i, j).text = celda
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def subido(nombre, contenido):
    return SimpleUploadedFile(nombre, contenido)


# ── Extracción de texto ───────────────────────────────────────────────

class TestExtraerTexto:
    def test_texto_plano_en_utf8_y_en_windows_1252(self):
        assert texto.extract_text(io.BytesIO('Semana 1: álgebra'.encode('utf-8')), 'a.txt') == 'Semana 1: álgebra'
        assert texto.extract_text(io.BytesIO('Semana 1: álgebra'.encode('cp1252')), 'a.TXT') == 'Semana 1: álgebra'

    def test_pdf_con_texto(self):
        contenido = pdf_con(['Semana 1', 'Numeros reales'], ['Semana 2', 'Funciones'])
        resultado = texto.extract_text(io.BytesIO(contenido), 'temario.pdf')
        assert 'Semana 1' in resultado and 'Numeros reales' in resultado and 'Funciones' in resultado

    def test_pdf_sin_texto_devuelve_vacio_no_error(self):
        assert texto.extract_text(io.BytesIO(pdf_con([])), 'escaneo.pdf') == ''

    def test_docx_con_parrafos_y_tablas(self):
        contenido = docx_con('Semana 1: Reales', tabla=[['Tema', 'Fecha'], ['Intervalos', '24 ago']])
        resultado = texto.extract_text(io.BytesIO(contenido), 'temario.docx')
        assert 'Semana 1: Reales' in resultado and 'Intervalos | 24 ago' in resultado

    @pytest.mark.parametrize('nombre', ['roto.pdf', 'roto.docx'])
    def test_un_archivo_dañado_lanza_error_controlado(self, nombre):
        with pytest.raises(texto.ExtractionError):
            texto.extract_text(io.BytesIO(b'esto no es un archivo valido'), nombre)

    def test_una_extension_desconocida_se_rechaza(self):
        with pytest.raises(texto.ExtractionError):
            texto.extract_text(io.BytesIO(b'x'), 'macro.exe')

    def test_limpia_nulos_espacios_y_saltos_de_mas_y_recorta(self):
        sucio = 'a\x00  b\t\tc\r\n\r\n\r\n\r\nd e'
        assert texto.clean_text(sucio) == 'a b c\n\nd e'
        assert len(texto.clean_text('x' * (texto.MAX_CHARS + 500))) == texto.MAX_CHARS


# ── Validación del archivo ────────────────────────────────────────────

class TestValidarArchivo:
    @pytest.mark.parametrize('nombre', ['t.pdf', 't.DOCX', 't.txt'])
    def test_acepta_pdf_word_y_texto(self, nombre):
        validate_syllabus_file(subido(nombre, b'x'))

    def test_el_doc_antiguo_explica_que_hacer(self):
        with pytest.raises(ValidationError) as error:
            validate_syllabus_file(subido('t.doc', b'x'))
        assert 'docx' in str(error.value)

    @pytest.mark.parametrize('nombre', ['t.html', 't.svg', 't.exe', 't.zip', 'sin_extension'])
    def test_rechaza_lo_que_un_navegador_podria_ejecutar_y_lo_desconocido(self, nombre):
        with pytest.raises(ValidationError):
            validate_syllabus_file(subido(nombre, b'x'))

    def test_rechaza_mas_de_10_mb(self):
        archivo = subido('t.pdf', b'x')
        archivo.size = MAX_SYLLABUS_BYTES + 1
        with pytest.raises(ValidationError):
            validate_syllabus_file(archivo)


# ── Guardar y extraer ─────────────────────────────────────────────────

class TestGuardarYExtraer:
    def test_subir_un_archivo_lo_guarda_en_privado_y_lo_deja_pendiente(self):
        curso = CourseFactory()
        temario = extraction.save_uploaded_file(curso, subido('Mi temario.pdf', pdf_con(['Semana 1'])))
        assert temario.status == Syllabus.Status.PENDING and temario.original_name == 'Mi temario.pdf'
        assert 'Mi temario' not in temario.file.name  # el nombre en disco no es el del docente
        assert temario.file.name.startswith(f'{curso.organization_id}/{curso.pk}/')
        assert temario.file.storage.exists(temario.file.name)

    def test_la_tarea_extrae_el_texto_y_marca_listo(self):
        temario = extraction.save_uploaded_file(CourseFactory(), subido('t.pdf', pdf_con(['Semana 1: Reales'])))
        assert extract_syllabus_text(temario.pk) == Syllabus.Status.READY
        temario.refresh_from_db()
        assert 'Semana 1: Reales' in temario.text and temario.status == Syllabus.Status.READY and temario.word_count >= 3

    def test_un_escaneo_queda_sin_texto_legible(self):
        temario = extraction.save_uploaded_file(CourseFactory(), subido('scan.pdf', pdf_con([])))
        assert extract_syllabus_text(temario.pk) == Syllabus.Status.EMPTY

    def test_un_archivo_dañado_queda_como_fallido_sin_romper(self):
        temario = extraction.save_uploaded_file(CourseFactory(), subido('roto.pdf', b'basura'))
        assert extract_syllabus_text(temario.pk) == Syllabus.Status.FAILED
        temario.refresh_from_db()
        assert temario.text == ''

    def test_la_tarea_de_un_temario_que_ya_no_existe_no_falla(self):
        assert extract_syllabus_text(999999) is None

    def test_reemplazar_el_archivo_borra_el_anterior_y_vuelve_a_extraer(self):
        curso = CourseFactory()
        primero = extraction.save_uploaded_file(curso, subido('a.txt', b'Semana 1: A'))
        ruta_vieja = primero.file.name
        segundo = extraction.save_uploaded_file(curso, subido('b.txt', b'Semana 1: B'))
        assert segundo.pk == primero.pk and Syllabus.objects.filter(course=curso).count() == 1
        assert not segundo.file.storage.exists(ruta_vieja)
        assert segundo.text == '' and segundo.status == Syllabus.Status.PENDING

    def test_pegar_texto_guarda_listo_y_quita_el_archivo(self):
        curso = CourseFactory()
        con_archivo = extraction.save_uploaded_file(curso, subido('a.txt', b'hola'))
        ruta = con_archivo.file.name
        pegado = extraction.save_pasted_text(curso, '  Semana 1: Reales  ')
        assert pegado.text == 'Semana 1: Reales' and pegado.status == Syllabus.Status.READY
        assert not pegado.file and pegado.original_name == '' and not pegado.file.storage.exists(ruta)

    def test_pegar_texto_vacio_lo_deja_sin_texto(self):
        assert extraction.save_pasted_text(CourseFactory(), '   ').status == Syllabus.Status.EMPTY

    def test_borrar_el_temario_o_el_grupo_borra_el_archivo(self):
        temario = extraction.save_uploaded_file(CourseFactory(), subido('a.txt', b'hola'))
        ruta, almacen = temario.file.name, temario.file.storage
        temario.course.delete()
        assert not Syllabus.objects.exists() and not almacen.exists(ruta)

    def test_la_organizacion_se_deriva_del_grupo(self):
        curso = CourseFactory()
        temario = SyllabusFactory(course=curso)
        semana = SyllabusWeekFactory(syllabus=temario)
        assert temario.organization_id == curso.organization_id == semana.organization_id


# ── Semanas ───────────────────────────────────────────────────────────

def temario_con(n, **kwargs):
    temario = SyllabusFactory(**kwargs)
    for i in range(1, n + 1):
        SyllabusWeekFactory(syllabus=temario, position=i, title=f'S{i}', topics='')
    return temario


def titulos(temario):
    return [(w.position, w.title) for w in temario.weeks.all()]


class TestSemanas:
    def test_agregar_va_al_final_con_el_numero_siguiente(self):
        temario = temario_con(2)
        nueva = semanas.add_week(temario, title='Nueva')
        assert nueva.position == 3 and temario.weeks.count() == 3

    def test_hay_un_tope_de_semanas(self):
        temario = temario_con(0)
        for _ in range(semanas.MAX_WEEKS):
            semanas.add_week(temario)
        with pytest.raises(semanas.WeekError):
            semanas.add_week(temario)

    def test_subir_y_bajar_intercambian_con_la_vecina(self):
        temario = temario_con(3)
        dos = temario.weeks.get(position=2)
        semanas.move_week(dos, -1)
        assert titulos(temario) == [(1, 'S2'), (2, 'S1'), (3, 'S3')]
        semanas.move_week(temario.weeks.get(title='S2'), +1)
        assert titulos(temario) == [(1, 'S1'), (2, 'S2'), (3, 'S3')]

    def test_en_los_extremos_mover_no_hace_nada(self):
        temario = temario_con(2)
        semanas.move_week(temario.weeks.get(position=1), -1)
        semanas.move_week(temario.weeks.get(position=2), +1)
        assert titulos(temario) == [(1, 'S1'), (2, 'S2')]

    def test_dividir_reparte_los_temas_y_recorre_las_siguientes(self):
        temario = temario_con(2)
        primera = temario.weeks.get(position=1)
        primera.topics = 'A\nB\nC\nD\nE'
        primera.save()
        nueva = semanas.split_week(primera)
        primera.refresh_from_db()
        assert primera.topic_list == ['A', 'B', 'C'] and nueva.topic_list == ['D', 'E']
        assert nueva.position == 2 and nueva.title == 'S1 (continuación)'
        assert titulos(temario) == [(1, 'S1'), (2, 'S1 (continuación)'), (3, 'S2')]

    def test_dividir_pide_al_menos_dos_temas(self):
        temario = temario_con(1)
        semana = temario.weeks.get()
        semana.topics = 'Solo uno'
        semana.save()
        with pytest.raises(semanas.WeekError):
            semanas.split_week(semana)
        assert temario.weeks.count() == 1

    def test_unir_junta_temas_objetivos_y_actividades_y_quita_la_siguiente(self):
        temario = temario_con(3)
        a, b = temario.weeks.get(position=1), temario.weeks.get(position=2)
        a.topics, a.objectives, a.assessment, b.topics, b.objectives, b.assessment = 'A', 'Oa', 'Tarea 1', 'B', 'Ob', 'Quiz 2'
        a.save(), b.save()
        unida = semanas.merge_with_next(a)
        assert unida.topic_list == ['A', 'B'] and unida.objectives == 'Oa\nOb' and unida.assessment == 'Tarea 1\nQuiz 2'
        assert titulos(temario) == [(1, 'S1'), (2, 'S3')]

    def test_unir_la_ultima_semana_no_se_puede(self):
        temario = temario_con(2)
        with pytest.raises(semanas.WeekError):
            semanas.merge_with_next(temario.weeks.get(position=2))

    def test_unir_conserva_sin_clase_solo_si_ambas_lo_eran(self):
        temario = temario_con(3)
        a, b, c = temario.weeks.order_by('position')
        a.no_class = b.no_class = True
        a.save(), b.save()
        assert semanas.merge_with_next(a).no_class is True
        a.refresh_from_db()
        assert semanas.merge_with_next(a).no_class is False  # la siguiente (c) sí tenía clase

    def test_quitar_renumera_sin_huecos(self):
        temario = temario_con(4)
        semanas.delete_week(temario.weeks.get(position=2))
        assert titulos(temario) == [(1, 'S1'), (2, 'S3'), (3, 'S4')]

    def test_renumerar_arregla_huecos_y_repetidos(self):
        temario = temario_con(0)
        for posicion in (5, 5, 9):
            SyllabusWeekFactory(syllabus=temario, position=posicion)
        semanas.renumber(temario)
        assert [w.position for w in temario.weeks.all()] == [1, 2, 3]


# ── Detectar semanas ──────────────────────────────────────────────────

TEMARIO = """Universidad Demo — Matemáticas II
Docente: Ana Pérez

Semana 1. Números reales y desigualdades
Temas: propiedades de los reales, intervalos; desigualdades lineales.
Objetivo: resolver desigualdades.

SEMANA 2: Funciones
- Dominio y rango
- Gráfica de una función
Objetivos: graficar funciones simples.
Interpretar tablas

Semana 3
Temas:
• Límites
• Continuidad
"""


class TestDetectarSemanas:
    def test_encuentra_las_semanas_con_titulo_temas_y_objetivos(self):
        resultado = semanas.detect_weeks(TEMARIO)
        assert [r['title'] for r in resultado] == ['Números reales y desigualdades', 'Funciones', '']
        assert resultado[0]['topics'].splitlines() == ['propiedades de los reales', 'intervalos', 'desigualdades lineales']
        assert resultado[0]['objectives'] == 'resolver desigualdades.'
        assert resultado[1]['topics'].splitlines() == ['Dominio y rango', 'Gráfica de una función']
        assert resultado[1]['objectives'].splitlines() == ['graficar funciones simples.', 'Interpretar tablas']
        assert resultado[2]['topics'].splitlines() == ['Límites', 'Continuidad']

    def test_lo_anterior_al_primer_encabezado_se_ignora(self):
        assert 'Docente' not in str(semanas.detect_weeks(TEMARIO))

    def test_sin_encabezados_no_propone_nada(self):
        assert semanas.detect_weeks('Un temario cualquiera\nsin formato\nUnidad 1: algo') == []
        assert semanas.detect_weeks('') == []

    def test_no_confunde_semana_dentro_de_una_frase(self):
        assert semanas.detect_weeks('Esta semana 3 veremos límites') == []

    def test_respeta_el_tope_de_semanas(self):
        texto_largo = '\n'.join(f'Semana {i}: T{i}' for i in range(1, 200))
        assert len(semanas.detect_weeks(texto_largo)) == semanas.MAX_WEEKS

    def test_aplicar_reemplaza_las_semanas_y_numera_en_orden(self):
        temario = temario_con(5)
        semanas.replace_weeks(temario, semanas.detect_weeks(TEMARIO))
        assert titulos(temario) == [(1, 'Números reales y desigualdades'), (2, 'Funciones'), (3, '')]


# Formato real de un temario en PDF: resumen inicial, «SEMANA N» con el título en la línea de abajo, secciones sin
# dos puntos, sesiones en tabla, encabezado y «Página N» en cada hoja, y secciones de cierre tras la última semana.
TEMARIO_PDF = """Matemáticas II – Temario de 4 semanas
Página 1
Matemáticas II
Duración
4 semanas · 3 sesiones por semana
Resumen del plan
Semana
Tema central
Cierre de la semana
1
Antiderivadas e integral indefinida
Quiz 1
2
Integral definida y áreas
Examen parcial

Matemáticas II – Temario de 4 semanas
Página 2
SEMANA 1
Antiderivadas e integral indefinida
Objetivo de la semana
Comprender la integración como proceso inverso de la derivación y resolver integrales
básicas y por cambio de variable.
Sesión
Contenidos
Sesión 1
Concepto de antiderivada y constante de integración. Relación entre
derivada e integral.
Sesión 2
Fórmulas básicas: potencias y exponenciales. Propiedades
de linealidad.
Actividades de aprendizaje
 Hoja de ejercicios de integrales inmediatas.
 Resolver 10 integrales por sustitución.
Entregable / evaluación
Tarea 1 y Quiz 1 en la Sesión 2.
SEMANA 2
Integral definida y áreas
Objetivo de la semana
Interpretar la integral definida como área acumulada y evaluarla con el Teorema Fundamental del
Cálculo.
Sesión
Contenidos
Sesión 1
Aproximación de áreas con rectángulos. Sumas de Riemann.

Matemáticas II – Temario de 4 semanas
Página 3
Sesión 2
Propiedades de la integral definida. Aplicación repetida.
Entregable / evaluación
Tarea 2 y Examen parcial.

Matemáticas II – Temario de 4 semanas
Página 4
Evaluación sugerida
Componente
Peso
Tareas semanales
20%
Bibliografía recomendada
 Stewart, J. Cálculo de una variable.
"""


class TestDetectarFormatoReal:
    @pytest.fixture
    def detectadas(self):
        return semanas.detect_weeks(TEMARIO_PDF)

    def test_encuentra_las_dos_semanas_y_no_toma_el_resumen_inicial_por_una(self, detectadas):
        assert len(detectadas) == 2

    def test_el_titulo_es_la_linea_de_abajo_del_encabezado(self, detectadas):
        assert [d['title'] for d in detectadas] == ['Antiderivadas e integral indefinida', 'Integral definida y áreas']

    def test_el_objetivo_junta_las_lineas_que_el_pdf_corto(self, detectadas):
        assert detectadas[0]['objectives'] == (
            'Comprender la integración como proceso inverso de la derivación y resolver integrales básicas y por cambio de variable.'
        )
        assert detectadas[1]['objectives'].endswith('Teorema Fundamental del Cálculo.')  # «Cálculo.» no queda suelto

    def test_cada_sesion_es_un_tema_con_su_contenido_completo(self, detectadas):
        assert detectadas[0]['topics'].splitlines() == [
            'Sesión 1: Concepto de antiderivada y constante de integración. Relación entre derivada e integral.',
            'Sesión 2: Fórmulas básicas: potencias y exponenciales. Propiedades de linealidad.',
        ]

    def test_las_actividades_y_los_entregables_van_a_actividades_y_evaluacion(self, detectadas):
        assert detectadas[0]['assessment'].splitlines() == [
            'Hoja de ejercicios de integrales inmediatas.', 'Resolver 10 integrales por sustitución.', 'Tarea 1 y Quiz 1 en la Sesión 2.',
        ]
        assert detectadas[1]['assessment'] == 'Tarea 2 y Examen parcial.'

    def test_una_etiqueta_con_contenido_en_la_misma_linea_tambien_cuenta(self):
        lineas = ['SEMANA 1', 'Reales', 'Tarea: leer el capítulo 3', 'Evaluación - Quiz corto', 'Recursos', 'Libro de texto']
        detectadas = semanas.detect_weeks('\n'.join(lineas))
        assert detectadas[0]['assessment'].splitlines() == ['leer el capítulo 3', 'Quiz corto']
        assert 'Libro de texto' not in str(detectadas)  # los recursos se siguen saltando

    def test_una_semana_sin_actividades_queda_con_el_campo_vacio(self):
        detectadas = semanas.detect_weeks('SEMANA 1\nReales\nSesión 1\nIntervalos.')
        assert detectadas[0]['assessment'] == ''

    def test_aplicar_guarda_las_actividades_y_evaluacion_en_cada_semana(self):
        temario = temario_con(1)
        semanas.replace_weeks(temario, semanas.detect_weeks(TEMARIO_PDF))
        assert temario.weeks.get(position=2).assessment == 'Tarea 2 y Examen parcial.'

    def test_actividades_y_entregables_no_se_mezclan_con_los_temas(self, detectadas):
        todo = ' '.join(d['topics'] + d['objectives'] for d in detectadas)
        for ajeno in ('Hoja de ejercicios', 'Resolver 10 integrales', 'Tarea 1', 'Quiz 1 en la Sesión'):
            assert ajeno not in todo

    def test_el_encabezado_y_el_numero_de_pagina_no_se_cuelan(self, detectadas):
        todo = ' '.join(d['title'] + d['topics'] + d['objectives'] for d in detectadas)
        assert 'Página' not in todo and 'Temario de 4 semanas' not in todo

    def test_una_sesion_que_cruza_de_pagina_sigue_en_su_semana(self, detectadas):
        assert detectadas[1]['topics'].splitlines() == [
            'Sesión 1: Aproximación de áreas con rectángulos. Sumas de Riemann.',
            'Sesión 2: Propiedades de la integral definida. Aplicación repetida.',
        ]

    def test_la_ultima_semana_no_se_traga_la_evaluacion_ni_la_bibliografia(self, detectadas):
        todo = detectadas[1]['topics'] + detectadas[1]['objectives']
        for cierre in ('Evaluación sugerida', 'Tareas semanales', 'Bibliografía', 'Stewart'):
            assert cierre not in todo

    def test_cortar_en_el_cierre_aunque_la_ultima_seccion_de_la_semana_sea_de_temas(self):
        lineas = ['SEMANA 1', 'Reales', 'Sesión 1', 'Intervalos y desigualdades.', 'Evaluación sugerida',
                  'Tareas semanales 20%', 'Bibliografía recomendada', 'Stewart']
        detectadas = semanas.detect_weeks('\n'.join(lineas))
        assert detectadas[0]['topics'] == 'Sesión 1: Intervalos y desigualdades.'

    def test_una_sesion_junta_todas_sus_lineas_aunque_una_termine_en_punto(self):
        lineas = ['SEMANA 1', 'Reales', 'Sesión 1', 'Primera frase del contenido.', 'Segunda frase.', 'Sesión 2', 'Otra.']
        detectadas = semanas.detect_weeks('\n'.join(lineas))
        assert detectadas[0]['topics'].splitlines() == ['Sesión 1: Primera frase del contenido. Segunda frase.', 'Sesión 2: Otra.']

    def test_con_saltos_de_linea_de_windows_da_lo_mismo(self):
        assert semanas.detect_weeks(TEMARIO_PDF.replace('\n', '\r\n')) == semanas.detect_weeks(TEMARIO_PDF)

    def test_un_parrafo_largo_tras_el_encabezado_no_se_toma_por_titulo(self):
        largo = 'Esta semana revisaremos con calma todos los conceptos del bloque y resolveremos dudas de la clase anterior, ' * 2
        detectadas = semanas.detect_weeks(f'Semana 1\n{largo.strip()}\nTemas: A, B')
        assert detectadas[0]['title'] == '' and 'revisaremos' in detectadas[0]['topics']

    def test_una_linea_que_empieza_con_semana_pero_es_prosa_larga_no_es_encabezado(self):
        assert semanas.detect_weeks('Semana 3 ' + 'palabra ' * 30) == []

    def test_varias_lineas_de_temas_sin_punto_se_quedan_separadas(self):
        detectadas = semanas.detect_weeks('Semana 1: X\nTemas:\nDominio\nRango\nGráfica')
        assert detectadas[0]['topics'].splitlines() == ['Dominio', 'Rango', 'Gráfica']

    def test_un_titulo_con_barra_o_guion_se_lee_igual(self):
        detectadas = semanas.detect_weeks('Semana 1 | Reales\nSemana 2 - Funciones\nSemana 3 – Límites')
        assert [d['title'] for d in detectadas] == ['Reales', 'Funciones', 'Límites']


# ── Quitar todas las semanas ──────────────────────────────────────────

class TestQuitarTodas:
    def test_quita_todas_y_dice_cuantas_eran(self):
        temario = temario_con(5)
        assert semanas.delete_all_weeks(temario) == 5 and temario.weeks.count() == 0

    def test_sin_semanas_no_hace_nada(self):
        assert semanas.delete_all_weeks(temario_con(0)) == 0

    def test_no_toca_las_semanas_de_otros_temarios(self):
        mio, otro = temario_con(2), temario_con(3)
        semanas.delete_all_weeks(mio)
        assert otro.weeks.count() == 3


# ── Fechas ────────────────────────────────────────────────────────────

class TestFechas:
    def test_sin_fecha_de_inicio_no_hay_fechas(self):
        temario = temario_con(1)
        assert semanas.week_dates(temario, temario.weeks.get()) is None
        assert semanas.current_week(temario, datetime.date(2026, 10, 20)) is None

    def test_cada_semana_va_de_lunes_a_viernes_aunque_el_inicio_sea_otro_dia(self):
        temario = temario_con(3, start_date=datetime.date(2026, 8, 26))  # miércoles
        lunes_a_viernes = [semanas.week_dates(temario, w) for w in temario.weeks.all()]
        assert lunes_a_viernes[0] == (datetime.date(2026, 8, 24), datetime.date(2026, 8, 28))
        assert lunes_a_viernes[1] == (datetime.date(2026, 8, 31), datetime.date(2026, 9, 4))
        assert lunes_a_viernes[2][0] == datetime.date(2026, 9, 7)

    def test_la_semana_actual_incluye_el_fin_de_semana_y_no_inventa_fuera_de_rango(self):
        temario = temario_con(3, start_date=datetime.date(2026, 8, 24))
        assert semanas.current_week(temario, datetime.date(2026, 8, 24)).position == 1
        assert semanas.current_week(temario, datetime.date(2026, 8, 30)).position == 1  # domingo
        assert semanas.current_week(temario, datetime.date(2026, 8, 31)).position == 2
        assert semanas.current_week(temario, datetime.date(2026, 8, 23)) is None
        assert semanas.current_week(temario, datetime.date(2026, 9, 14)) is None


def test_un_grupo_tiene_un_solo_temario():
    curso = CourseFactory()
    SyllabusFactory(course=curso)
    with pytest.raises(Exception):
        Syllabus.objects.create(course=curso)


def test_las_semanas_se_ordenan_por_numero():
    temario = SyllabusFactory()
    SyllabusWeekFactory(syllabus=temario, position=2, title='B')
    SyllabusWeekFactory(syllabus=temario, position=1, title='A')
    assert [w.title for w in SyllabusWeek.objects.filter(syllabus=temario)] == ['A', 'B']
